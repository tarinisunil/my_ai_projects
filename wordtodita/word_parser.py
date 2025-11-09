from docx import Document

def parse_word_doc(path: str):
    """Parse .docx and return structured elements"""
    doc = Document(path)
    content = []

    for para in doc.paragraphs:
        if para.style.name.startswith("Heading"):
            level = int(para.style.name[-1]) if para.style.name[-1].isdigit() else 1
            content.append({"type": "heading", "level": level, "text": para.text})
        elif para.text.strip():
            content.append({"type": "paragraph", "text": para.text})
    
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        content.append({"type": "table", "rows": rows})

    return content
